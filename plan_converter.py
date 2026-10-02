# -*- coding: utf-8 -*-
"""제공계획서 변환 엔진 (도형이의 제공계획서 변환기 1.0.3 에서 화면 없이 옮겨 온 것. 예전 변환기와는 따로 운영).

- 바탕화면의 '제공계획서 작업전' 엑셀을 값만 추출·검사해 '작업후' 폴더에 이름으로 저장하고 통합본을 만든다.
- 대상자 교차검증은 맞춤돌봄도우미가 하루 한 번 받는 대상자리스트(C:\\맞춤돌봄도우미)로,
  생활지원사 생년월일 검증은 goodeos 에서 받은 지원사 명단으로 한다.
"""
import calendar
import os
import re
import zipfile
import io
import tempfile
import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xlrd
except ImportError:
    xlrd = None

SOURCE_FOLDER_NAME = "제공계획서 작업전"
TARGET_FOLDER_NAME = "제공계획서 작업후"
TITLE_KEYWORD = "제공"
TITLE_KEYWORDS = (TITLE_KEYWORD, "계획서")

def desktop_dir() -> Path:
    """윈도우가 알려 주는 진짜 바탕화면 위치 (OneDrive 바탕화면이어도 맞게 나옴)."""
    try:
        import ctypes
        from ctypes import wintypes

        class GUID(ctypes.Structure):
            _fields_ = [("Data1", wintypes.DWORD), ("Data2", wintypes.WORD), ("Data3", wintypes.WORD),
                        ("Data4", ctypes.c_ubyte * 8)]

        FOLDERID_Desktop = GUID(0xB4BFCC3A, 0xDB2C, 0x424C,
                                (ctypes.c_ubyte * 8)(0xB0, 0x29, 0x7F, 0xE9, 0x9A, 0x87, 0xC6, 0x41))
        buf = ctypes.c_wchar_p()
        if ctypes.windll.shell32.SHGetKnownFolderPath(ctypes.byref(FOLDERID_Desktop), 0, None,
                                                       ctypes.byref(buf)) == 0 and buf.value:
            path = Path(buf.value)
            ctypes.windll.ole32.CoTaskMemFree(buf)
            if path.is_dir():
                return path
    except Exception:
        pass
    home = Path.home()
    for p in (home / "OneDrive" / "바탕 화면", home / "OneDrive" / "Desktop", home / "바탕 화면", home / "Desktop"):
        if p.is_dir():
            return p
    return home


def work_folders():
    """바탕화면의 '제공계획서 작업전/작업후' 폴더. 없으면 만든다. → (작업전, 작업후)"""
    base = desktop_dir()
    src, dst = base / SOURCE_FOLDER_NAME, base / TARGET_FOLDER_NAME
    src.mkdir(parents=True, exist_ok=True)
    dst.mkdir(parents=True, exist_ok=True)
    return src, dst


def _natural_sort_key(path: Path):
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", path.stem)]


def find_source_excels(folder):
    patterns = ["*.xlsx", "*.xlsm", "*.xls"]
    matches = []
    for pattern in patterns:
        matches.extend(folder.glob(pattern))

    matches = [
        f for f in matches
        if any(k in f.stem for k in TITLE_KEYWORDS) and not f.name.startswith("~$")
    ]

    if not matches:
        raise FileNotFoundError(f"'{folder}' '{'/'.join(TITLE_KEYWORDS)}' 엑셀 파일을 찾지 못했습니다.")

    matches.sort(key=_natural_sort_key)
    return matches


def _extract_person_name(stem: str) -> str:
    match = re.search(r"[가-힣]{2,}", stem)
    if match:
        name = match.group(0)
        if not any(k in name for k in TITLE_KEYWORDS):
            return name
    return stem


def _unique_target_path(target_folder: Path, base_name: str) -> Path:
    candidate = target_folder / f"{base_name}.xlsx"
    n = 1
    while candidate.exists():
        candidate = target_folder / f"{base_name} ({n}).xlsx"
        n += 1
    return candidate


def _strip_alternate_content(xml_text: str) -> str:
    pattern = re.compile(
        r"<(?:\w+:)?AlternateContent\b[^>]*>.*?<(?:\w+:)?Fallback\s*>(.*?)</(?:\w+:)?Fallback>\s*</(?:\w+:)?AlternateContent>",
        re.DOTALL,
    )
    return pattern.sub(r"\1", xml_text)


def _patch_missing_style_names(xml_bytes: bytes) -> bytes:
    text = xml_bytes.decode("utf-8", errors="ignore")

    def add_name(match):
        tag = match.group(0)
        if re.search(r"\bname\s*=", tag):
            return tag
        if tag.endswith("/>"):
            return tag[:-2].rstrip() + ' name="Normal"/>'
        return tag[:-1].rstrip() + ' name="Normal">'

    text = re.sub(r"<(?:\w+:)?cellStyle(?![a-zA-Z])[^>]*/?>", add_name, text)
    return text.encode("utf-8")


def _max_style_index_used(zin, names_in_zip):
    max_s = -1
    for name in names_in_zip:
        if not re.match(r"xl/worksheets/sheet\d+\.xml$", name):
            continue
        data = zin.read(name).decode("utf-8", errors="ignore")
        for m in re.finditer(r'\ss="(\d+)"', data):
            v = int(m.group(1))
            if v > max_s:
                max_s = v
    return max_s


def _max_xfid_referenced(xml_text):
    max_id = -1
    for m in re.finditer(r'xfId="(\d+)"', xml_text):
        v = int(m.group(1))
        if v > max_id:
            max_id = v
    return max_id


def _pad_xf_list(xml_text, tag_name, needed_count):
    pattern = re.compile(
        "<((?:\\w+:)?" + tag_name + ')\\s+count="(\\d+)"\\s*>(.*?)</(?:\\w+:)?' + tag_name + ">",
        re.DOTALL,
    )

    match = pattern.search(xml_text)
    if not match:
        return xml_text

    full_open_name, _count_str, body = match.groups()
    prefix = ""
    if ":" in full_open_name:
        prefix = full_open_name.split(":")[0] + ":"

    xf_tag = prefix + "xf"

    xf_items = re.findall(
        "<" + re.escape(xf_tag) + r"\b[^>]*?/>|<" + re.escape(xf_tag) + r"\b[^>]*?>.*?</" + re.escape(xf_tag) + ">",
        body,
        re.DOTALL,
    )

    current_count = len(xf_items)
    if needed_count <= current_count:
        return xml_text

    if current_count == 0:
        last_item = f'<{xf_tag} numFmtId="0" fontId="0" fillId="0" borderId="0"/>'
    else:
        last_item = xf_items[-1]

    extra = last_item * (needed_count - current_count)
    new_body = body + extra
    new_section = f'<{full_open_name} count="{needed_count}">{new_body}</{full_open_name}>'

    return xml_text[:match.start()] + new_section + xml_text[match.end():]


def _read_bytes_io(path):
    return io.BytesIO(Path(path).read_bytes())


def _try_patched_load(path, tmp_path):
    with zipfile.ZipFile(_read_bytes_io(path), "r") as zin:
        names_in_zip = zin.namelist()
        max_s = _max_style_index_used(zin, names_in_zip)

        with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in names_in_zip:
                data = zin.read(name)
                if name == "xl/styles.xml":
                    text = data.decode("utf-8", errors="ignore")
                    text = _strip_alternate_content(text)
                    data = _patch_missing_style_names(text.encode("utf-8"))
                    text = data.decode("utf-8", errors="ignore")
                    max_xfid = _max_xfid_referenced(text)
                    if max_s >= 0:
                        text = _pad_xf_list(text, "cellXfs", max_s + 1)
                    if max_xfid >= 0:
                        text = _pad_xf_list(text, "cellStyleXfs", max_xfid + 1)
                    data = text.encode("utf-8")
                zout.writestr(name, data)

    return openpyxl.load_workbook(_read_bytes_io(tmp_path), data_only=True)


MINIMAL_STYLES_XML = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>
  <fills count="2">
    <fill><patternFill patternType="none"/></fill>
    <fill><patternFill patternType="gray125"/></fill>
  </fills>
  <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
  <cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs>
  <cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>
</styleSheet>"""


def _try_stripped_load(path, tmp_path):
    with zipfile.ZipFile(_read_bytes_io(path), "r") as zin:
        with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in zin.namelist():
                data = zin.read(name)
                if name == "xl/styles.xml":
                    data = MINIMAL_STYLES_XML
                elif re.match(r"xl/worksheets/sheet\d+\.xml$", name):
                    text = data.decode("utf-8", errors="ignore")
                    text = re.sub(r'\ss="\d+"', "", text)
                    text = re.sub(r'\sstyle="\d+"', "", text)
                    text = re.sub(
                        r"<(?:\w+:)?conditionalFormatting\b.*?</(?:\w+:)?conditionalFormatting>",
                        "",
                        text,
                        flags=re.DOTALL,
                    )
                    text = re.sub(r"<(?:\w+:)?conditionalFormatting\b[^>]*/>", "", text)
                    data = text.encode("utf-8")
                zout.writestr(name, data)

    return openpyxl.load_workbook(_read_bytes_io(tmp_path), data_only=True)


class _XlsCellProxy:
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class _XlsSheetAdapter:

    def __init__(self, sheet, book):
        self._sheet = sheet
        self._book = book
        self.max_row = sheet.nrows
        self.max_column = sheet.ncols

    def cell(self, row, column):
        r = row - 1
        c = column - 1
        if r < 0 or r >= self._sheet.nrows or c < 0 or c >= self._sheet.ncols:
            return _XlsCellProxy(None)

        xl_cell = self._sheet.cell(r, c)
        value = xl_cell.value
        ctype = xl_cell.ctype

        if ctype == xlrd.XL_CELL_EMPTY or ctype == xlrd.XL_CELL_BLANK:
            value = None
        elif ctype == xlrd.XL_CELL_TEXT and value == "":
            value = None
        elif ctype == xlrd.XL_CELL_DATE:
            try:
                value = datetime.datetime(*xlrd.xldate_as_tuple(value, self._book.datemode))
            except Exception:
                pass
        elif ctype == xlrd.XL_CELL_BOOLEAN:
            value = bool(value)

        return _XlsCellProxy(value)


class _XlsWorkbookAdapter:

    def __init__(self, sheet, book):
        self.active = _XlsSheetAdapter(sheet, book)


def _load_xls_workbook(path):
    if xlrd is None:
        raise RuntimeError("'.xls' 파일을 읽으려면 xlrd 패키지가 필요합니다. (py -m pip install xlrd)")
    book = xlrd.open_workbook(file_contents=Path(path).read_bytes())
    sheet = book.sheet_by_index(0)
    return _XlsWorkbookAdapter(sheet, book)


def load_workbook_safely(path, log):
    if path.suffix.lower() == ".xls":
        return _load_xls_workbook(path)

    try:
        return openpyxl.load_workbook(_read_bytes_io(path), data_only=True)
    except Exception as first_err:
        for attempt_name, attempt_func in (
            ("서식 보정", _try_patched_load),
            ("서식 없이 값만 추출", _try_stripped_load),
        ):
            tmp_path = None
            try:
                tmp_fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
                os.close(tmp_fd)
                wb = attempt_func(path, tmp_path)
                log(f"    (자동 복구 성공: {attempt_name})")
                return wb
            except Exception:
                pass
            finally:
                if tmp_path and os.path.exists(tmp_path):
                    try:
                        os.remove(tmp_path)
                    except OSError:
                        pass
        raise first_err


SERVICE_CODES = {
    "11201": "방문-정보제공(사회.재난안전.보건.복지 정보제공)",
    "11301": "방문-생활안전점검(안전.위생관리.기타)",
    "11401": "방문-말벗(정서지원)",
    "11501": "방문-안전+안부+말벗",
    "11601": "방문-안전.안부+말벗+생활안전점검",
    "1A102": "방문-안전.안부+말벗+생활안전점검+정보제공",
    "11801": "방문-안전.안부+말벗+정보제공",
    "12101": "전화-안전.안부확인",
    "12201": "전화-정보제공(사회.재난안전.보건.복지 정보제공)",
    "12301": "전화-말벗(정서지원)",
    "12401": "전화-안전.안부+정보제공",
    "19101": "전화-안전.안부+정보제공+말벗",
    "12606": "전화-안전.안부+말벗",
    "13101": "ICT기기 사용방법교육",
    "13201": "ICT데이터 확인.점검",
    "14101": "ICT 관리.교육",
    "14102": "디지털 백신",
    "14201": "ICT 안전.안부확인",
    "21101": "문화여가활동",
    "21201": "평생교육활동",
    "21301": "체험여행활동",
    "21401": "여가활동",
    "21501": "문화활동",
    "22101": "자조모임",
    "31101": "영양교육",
    "33202": "보건교육",
    "33303": "건강교육",
    "31401": "영양+보건교육",
    "31501": "보건+건강교육",
    "31601": "영양+건강교육",
    "31701": "영양+보건+건강교육",
    "32101": "우울예방프로그램",
    "34202": "인지활동프로그램",
    "32301": "우울+인지프로그램",
    "41101": "외출동행",
    "43101": "식사관리",
    "43202": "청소관리",
    "42301": "식사+청소관리",
    "42401": "식사+청소+외출동행",
    "51101": "생활용품지원",
    "51201": "식료품지원",
    "51301": "후원금지원",
    "51401": "생활용품+식료품지원",
    "51501": "생활용품+식료품+후원금지원",
    "51601": "생활용품+기타서비스",
    "51701": "식료품지원+기타서비스",
    "51801": "생활용품+식료품+후원금지원+기타서비스",
    "52101": "주거위생개선지원",
    "52201": "주거환경개선지원",
    "52301": "주거위생+주거환경개선지원",
    "53101": "의료연계지원",
    "53201": "건강보조지원",
    "53301": "의료연계지원+건강보조지원",
    "54101": "기타서비스",
    "54201": "스마트빌리지",
    "55101": "영양교육+보건교육+건강교육",
    "55201": "우울예방+인지활동프로그램",
    "55301": "평생교육프로그램",
    "81101": "혹한기(겨울)",
    "81102": "혹서기(여름)",
    "81103": "기타",
    "91101": "광역실무협의체",
    "91102": "시군구실무협의체",
    "91103": "사례실무회의",
    "91104": "시군구심의판정위원회",
    "91201": "조회",
    "91202": "정기회의(월)",
    "91203": "정기회의(주)",
    "91204": "정기회의(기타)",
    "91205": "실무회의",
    "91206": "기타회의",
    "92101": "기초직무교육1",
    "92102": "기초직무교육2",
    "92103": "기초직무교육3",
    "92104": "역랑강화교육(중앙)",
    "92105": "역량강화교육(광역)",
    "92106": "심화교육",
    "92201": "법정의무교육",
    "92202": "기타교육",
    "93101": "선정조사,상담,제공계획",
    "93201": "기관업무지원",
    "93301": "지자체지원",
    "15101": "전화 안부확인",
    "15202": "방문 안부확인",
    "15303": "AI·디지털 안부확인",
    "16101": "생활안전점검·관리",
    "17101": "지지·격려·공감",
    "18101": "사회·재난대응 정보제공",
    "18202": "보건·복지 정보제공",
    "18303": "기타생활 정보제공",
    "18401": "사회·재난대응+보건·복지",
    "18501": "사회·재난대응+기타생활",
    "18601": "보건·복지+기타생활",
    "18701": "사회·재난대응+보건·복지+기타생활",
    "23101": "평생교육활동",
    "23202": "사회관계활동",
    "23303": "여가문화활동",
    "23404": "기타사회참여활동",
    "24101": "자조모임",
    "33101": "영양교육 프로그램",
    "33401": "영양+보건교육",
    "33501": "보건+건강교육",
    "33601": "영양+건강교육",
    "33701": "영양+보건+건강교육",
    "34101": "우울예방 프로그램",
    "34301": "우울+인지 프로그램",
    "43301": "식사+청소관리",
    "43401": "식사+청소+외출동행",
    "A1101": "개별상담·프로그램",
    "A1202": "집단상담·프로그램",
    "A1303": "치료지원",
    "B1101": "영양지원",
    "B1202": "가사지원",
    "B1303": "동행지원",
    "57101": "기타 서비스 연계",
    "92107": "기초교육",
    "92108": "기초+심화교육",
    "92109": "역량강화교육(지역)",
    "92110": "심리지원교육(중앙)",
    "92111": "심리지원교육(광역)",
    "92112": "심리지원교육(지역)",
    "92113": "직무교육",
    "C1101": "전화 안부확인",
    "C1202": "방문 안부확인",
    "C1303": "AI·디지털 안부확인",
    "C2101": "생활안전점검·관리",
    "C3101": "말벗(지지·격려·공감)",
    "C4101": "정보제공(사회·재난대응+보건·복지+기타생활)",
    "C5101": "전화안부+말벗+정보제공",
    "C6101": "방문안부+생활안전+말벗+정보제공",
    "C7101": "생활지원연계(생활용품+식료품+후원금+기타지원)",
    "C8101": "주거개선연계(주거위생+주거환경개선지원)",
    "C9101": "건강지원연계(의료연계+건강보조지원)",
    "CA101": "기타 서비스 연계",
}

STANDARD_HEADERS = [
    "서비스일자", "계획 시작시간", "계획 종료시간", "실적 시작시간", "실적 종료시간",
    "대상자명", "대상자 생년월일", "대상자 성별",
    "생활지원사명", "생활지원사 생년월일", "생활지원사 성별",
    "서비스코드", "거래처코드(자원)",
]
NUM_COLS = len(STANDARD_HEADERS)

COL_DATE = 0
COL_START = 1
COL_END = 2
COL_A_START = 3
COL_A_END = 4
COL_NAME = 5
COL_BIRTH = 6
COL_GENDER = 7
COL_WORKER = 8
COL_WORKER_BIRTH = 9
COL_WORKER_GENDER = 10
COL_CODE = 11
COL_VENDOR = 12

END_CHECK_SKIP_COLS = {COL_A_START, COL_A_END}

BENEFICIARY_FILE_STEM = "대상자리스트"
GROUP_COLUMN_NAME = "관리구분"  # 대상자리스트 '이용정보' 아래 칸, 값은 일반/중점
GENERAL_GROUP_FORBIDDEN_CODES = {"43202": "청소관리"}  # 일반군에 넣으면 안 되는 서비스코드

def _is_blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def _list_cell_to_text(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y%m%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def load_beneficiary_list(path, log=None):
    """대상자리스트 파일(goodeos 대상자조회 출력 원본) → {이름: [대상자 정보]}"""
    path = Path(path)
    raw = path.read_bytes()
    rows = []
    is_html = raw.lstrip()[:1] == b"<"
    html_trs = None
    if is_html:
        text = None
        for enc in ("utf-8", "cp949", "utf-16"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise RuntimeError(f"'{path.name}' 파일의 문자 인코딩을 알 수 없습니다.")
        html_trs = re.findall(r"<tr[^>]*>(.*?)</tr>", text, re.S | re.I)
        for tr in html_trs:
            cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, re.S | re.I)
            cells = [re.sub(r"<[^>]+>", "", c).replace("&nbsp;", " ").strip() for c in cells]
            rows.append(cells)
    elif path.suffix.lower() == ".xls":
        if xlrd is None:
            raise RuntimeError("대상자리스트(.xls)를 읽으려면 xlrd 패키지가 필요합니다.")
        book = xlrd.open_workbook(file_contents=path.read_bytes())
        sheet = book.sheet_by_index(0)
        for r in range(sheet.nrows):
            row_cells = []
            for c in range(sheet.ncols):
                v = sheet.cell_value(r, c)
                if sheet.cell_type(r, c) == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate_as_datetime(v, book.datemode)
                    except Exception:
                        pass
                row_cells.append(_list_cell_to_text(v))
            rows.append(row_cells)
    else:
        wb = openpyxl.load_workbook(_read_bytes_io(path), data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            rows.append([_list_cell_to_text(v) for v in row])

    header_row_i = None
    for i, cells in enumerate(rows[:5]):
        if "성명" in cells and "생년월일" in cells:
            header_row_i = i
            break
    if header_row_i is None:
        raise RuntimeError(f"'{path.name}'에서 성명/생년월일 열을 찾지 못했습니다.")

    col_idx = {}
    WANTED = ("성명", "생년월일", "성별", "생활지원사", "이용상태")
    if is_html:
        attr_cells = re.findall(r"<t[dh]([^>]*)>(.*?)</t[dh]>", html_trs[header_row_i], re.S | re.I)
        real_idx = 0
        for attrs, content in attr_cells:
            name = re.sub(r"<[^>]+>", "", content).replace("&nbsp;", " ").strip()
            m = re.search(r'colspan\s*=\s*["\']?(\d+)', attrs, re.I)
            span = int(m.group(1)) if m else 1
            if name in WANTED and name not in col_idx:
                col_idx[name] = real_idx
            real_idx += span
        # 2단 헤더: 윗줄에서 rowspan=2가 아닌 칸은 아랫줄 칸을 차례로 가짐 (예: 이용정보 → 관리구분)
        if header_row_i + 1 < len(html_trs):
            sub_names = rows[header_row_i + 1]
            sub_i = 0
            real_idx = 0
            for attrs, _content in attr_cells:
                m = re.search(r'colspan\s*=\s*["\']?(\d+)', attrs, re.I)
                span = int(m.group(1)) if m else 1
                if not re.search(r'rowspan\s*=\s*["\']?2', attrs, re.I):
                    for k in range(span):
                        if sub_i < len(sub_names):
                            if sub_names[sub_i] == GROUP_COLUMN_NAME and GROUP_COLUMN_NAME not in col_idx:
                                col_idx[GROUP_COLUMN_NAME] = real_idx + k
                            sub_i += 1
                real_idx += span
    else:
        header_cells = rows[header_row_i]
        for key in WANTED:
            if key in header_cells:
                col_idx[key] = header_cells.index(key)
        for cells in rows[header_row_i:header_row_i + 2]:
            if GROUP_COLUMN_NAME in cells and GROUP_COLUMN_NAME not in col_idx:
                col_idx[GROUP_COLUMN_NAME] = cells.index(GROUP_COLUMN_NAME)

    if "성명" not in col_idx:
        raise RuntimeError(f"'{path.name}'에서 성명 열의 실제 위치를 계산하지 못했습니다.")

    ACTIVE_STATUS_PREFIXES = ("이용", "장기부재")
    INACTIVE_STATUS_KEYWORDS = ("종료", "중지", "불가", "정지", "보류", "대기", "탈락", "사망")

    def _is_active_status(s):
        s2 = re.sub(r"\s+", "", s)
        if any(k in s2 for k in INACTIVE_STATUS_KEYWORDS):
            return False
        return any(s2.startswith(p) for p in ACTIVE_STATUS_PREFIXES)

    index = {}
    for cells in rows[header_row_i + 1:]:
        if len(cells) <= col_idx["성명"]:
            continue
        name = cells[col_idx["성명"]].strip()
        if not name or not re.match(r"^[가-힣]{2,}", name):
            continue
        birth_raw = cells[col_idx.get("생년월일", -1)].strip() if "생년월일" in col_idx and len(cells) > col_idx["생년월일"] else ""
        digits = re.sub(r"\D", "", birth_raw)
        if len(digits) != 8:
            continue
        birth = int(digits)
        if "이용상태" in col_idx:
            status = cells[col_idx["이용상태"]].strip() if len(cells) > col_idx["이용상태"] else ""
            if not _is_active_status(status):
                continue
        entry = {
            "name": name,
            "birth": birth,
            "gender": cells[col_idx["성별"]].strip() if "성별" in col_idx and len(cells) > col_idx["성별"] else "",
            "worker": cells[col_idx["생활지원사"]].strip() if "생활지원사" in col_idx and len(cells) > col_idx["생활지원사"] else "",
            "group": cells[col_idx[GROUP_COLUMN_NAME]].strip() if GROUP_COLUMN_NAME in col_idx and len(cells) > col_idx[GROUP_COLUMN_NAME] else "",
        }
        index.setdefault(name, []).append(entry)

    if not index:
        raise RuntimeError(
            f"'{path.name}'에서 이용상태가 '이용'/'장기부재'인 대상자를 한 명도 읽지 못했습니다. "
            f"공단 시스템에서 받은 원본 파일을 그대로(엑셀로 다시 저장하지 말고) 넣어주세요.")

    for name, entries in index.items():
        if len(entries) > 1:
            for idx, e in enumerate(entries):
                e["display_name"] = f"{name}{chr(ord('A') + idx)}" if idx < 26 else f"{name}{idx + 1}"
        else:
            entries[0]["display_name"] = name

    return index


def find_header_row(ws, max_scan=30):
    marker_set = set(STANDARD_HEADERS)
    limit = min(ws.max_row, max_scan)
    for r in range(1, limit + 1):
        hits = 0
        for c in range(1, NUM_COLS + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip() in marker_set:
                hits += 1
        if hits >= 4:
            return r
    return None


def extract_data_grid(ws, log):
    header_row = find_header_row(ws)
    if header_row is None:
        first = ws.cell(row=1, column=1).value
        if re.fullmatch(r"\d{8}", re.sub(r"\D", "", str(first or ""))):
            log("    (헤더 행이 없어 표준 헤더를 넣고 1행부터 데이터로 읽습니다)")
            data_start = 1
        else:
            data_start = 2
    else:
        if header_row > 1:
            log(f"    (헤더 위의 잘못된 {header_row - 1}행을 제거했습니다)")
        data_start = header_row + 1

    grid = []
    row_nums = []
    for r in range(data_start, ws.max_row + 1):
        values = [ws.cell(row=r, column=c).value for c in range(1, NUM_COLS + 1)]
        blanks = sum(
            1 for i, v in enumerate(values)
            if i not in END_CHECK_SKIP_COLS and _is_blank(v)
        )
        if blanks >= 2:
            break
        grid.append(values)
        row_nums.append(r)
    return grid, row_nums


def normalize_time_value(v):
    if v is None:
        return None
    if isinstance(v, datetime.time):
        return f"{v.hour:02d}{v.minute:02d}"
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, int):
        s = str(v)
    else:
        s = str(v).strip().replace(":", "")
        s = re.sub(r"[^\d]", "", s)
    if not s.isdigit() or len(s) == 0 or len(s) > 4:
        return None
    s = s.zfill(4)
    hh, mm = int(s[:2]), int(s[2:])
    if hh > 24 or mm > 59:
        return None
    return s


def time_to_minutes(s):
    return int(s[:2]) * 60 + int(s[2:])


def minutes_to_time(m):
    return f"{m // 60:02d}{m % 60:02d}"


def normalize_date_value(v):
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.year * 10000 + v.month * 100 + v.day
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    digits = re.sub(r"\D", "", str(v).strip())
    if len(digits) != 8:
        return None
    n = int(digits)
    mm, dd = (n // 100) % 100, n % 100
    if not (1 <= mm <= 12 and 1 <= dd <= 31):
        return None
    return n


def clean_service_code(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"[^0-9A-Za-z]", "", str(v)).upper()


DEFAULT_VENDOR_CODE = "0002"


def normalize_vendor_code(v):
    if _is_blank(v):
        return DEFAULT_VENDOR_CODE
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    if s.isdigit():
        return s.zfill(4)
    return s


def _display_value(v):
    if v is None:
        return "(빈칸)"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


class FixProposal:
    _next_id = [0]

    def __init__(self, kind, file_label, row_label, person, current_disp, suggested_disp,
                 reason, cells, value_builder):
        FixProposal._next_id[0] += 1
        self.pid = FixProposal._next_id[0]
        self.kind = kind
        self.file_label = file_label
        self.row_label = row_label
        self.person = person
        self.current_disp = current_disp
        self.suggested_disp = suggested_disp
        self.reason = reason
        self.cells = cells
        self.value_builder = value_builder

    def build_values(self, text):
        return self.value_builder(text)


def _single_time_builder(text):
    t = normalize_time_value(text)
    if t is None:
        raise ValueError("시간은 0900 같은 4자리로 입력하세요.")
    return [t]


def _time_range_builder(text):
    parts = re.split(r"[~\-]", str(text))
    if len(parts) != 2:
        raise ValueError("시간 범위는 0900~1000 형식으로 입력하세요.")
    a, b = normalize_time_value(parts[0]), normalize_time_value(parts[1])
    if a is None or b is None:
        raise ValueError("시간 범위는 0900~1000 형식으로 입력하세요.")
    return [a, b]


def _date_builder(text):
    d = normalize_date_value(text)
    if d is None:
        raise ValueError("날짜는 20260706 같은 8자리로 입력하세요.")
    return [d]


def _text_builder(text):
    return [str(text).strip()]


def _code_builder(text):
    return [clean_service_code(text)]


_DATE_SEP = r"[~\-ㅡ–—]"


def _valid_md(mm, dd):
    return 1 <= mm <= 12 and 1 <= dd <= 31


def parse_filename_dates(stem):
    matches = re.findall(r"(\d{1,2})\s*월\s*(\d{1,2})\s*일", stem)
    if matches:
        s = (int(matches[0][0]), int(matches[0][1]))
        if _valid_md(*s):
            if len(matches) >= 2:
                e = (int(matches[1][0]), int(matches[1][1]))
                if _valid_md(*e):
                    return s, e
            return s, None

    m = re.search(rf"(?<!\d)(\d{{4}})\s*{_DATE_SEP}\s*(\d{{4}})(?!\d)", stem)
    if m:
        s = (int(m.group(1)[:2]), int(m.group(1)[2:]))
        e = (int(m.group(2)[:2]), int(m.group(2)[2:]))
        if _valid_md(*s) and _valid_md(*e):
            return s, e

    m = re.search(
        rf"(?<!\d)(\d{{1,2}})[.,](\d{{1,2}})\s*{_DATE_SEP}\s*(?:(\d{{1,2}})[.,])?(\d{{1,2}})(?!\d)",
        stem)
    if m:
        s = (int(m.group(1)), int(m.group(2)))
        e = (int(m.group(3)) if m.group(3) else s[0], int(m.group(4)))
        if _valid_md(*s) and _valid_md(*e):
            return s, e

    m = re.search(rf"(?<!\d)(\d{{4}})\s*{_DATE_SEP}\s*(\d{{1,2}})(?!\d)", stem)
    if m:
        s = (int(m.group(1)[:2]), int(m.group(1)[2:]))
        e = (s[0], int(m.group(2)))
        if _valid_md(*s) and _valid_md(*e):
            return s, e

    for m in re.finditer(r"(?<!\d)(\d{4})(?!\d)", stem):
        s = (int(m.group(1)[:2]), int(m.group(1)[2:]))
        if _valid_md(*s):
            return s, None
    m = re.search(r"(?<!\d)(\d{1,2})[.,](\d{2})(?!\d)", stem)
    if m:
        s = (int(m.group(1)), int(m.group(2)))
        if _valid_md(*s):
            return s, None

    return None, None


def is_revision_file(stem):
    return "수정" in stem


class FileJob:

    def __init__(self, src_path):
        self.src_path = src_path
        self.grid = []
        self.row_nums = []
        self.person_name = _extract_person_name(src_path.stem)
        self.file_date_start, self.file_date_end = parse_filename_dates(src_path.stem)
        self.is_revision = is_revision_file(src_path.stem)
        self.proposals = []
        self.errors = []
        self.warnings = []
        self.auto_fixes = []
        self.load_error = None
        self.dropped_rows = {}

    @property
    def worker_name(self):
        for row in self.grid:
            v = row[COL_WORKER]
            if not _is_blank(v):
                return str(v).strip()
        return self.person_name


def _norm_str(v):
    return "" if v is None else str(v).strip()


def scan_file(job: FileJob, ws, beneficiaries, log, worker_data=None):
    label = job.src_path.name
    job.grid, job.row_nums = extract_data_grid(ws, log)
    if not job.grid:
        job.load_error = "데이터를 찾을 수 없음"
        return

    grid, row_nums = job.grid, job.row_nums

    if beneficiaries is not None:
        _cross_check_beneficiaries(job, beneficiaries)

    if worker_data:  # goodeos 에서 지원사 명단을 받았을 때만
        _cross_check_workers(job, worker_data)

    silent_count = 0
    for i, row in enumerate(grid):
        rlabel = row_nums[i]
        person = _norm_str(row[COL_NAME])

        for col, colname in ((COL_START, "계획 시작시간"), (COL_END, "계획 종료시간"),
                             (COL_A_START, "실적 시작시간"), (COL_A_END, "실적 종료시간")):
            v = row[col]
            if _is_blank(v):
                continue
            t = normalize_time_value(v)
            if t is None:
                job.warnings.append(f"{rlabel}행 {colname} '{_display_value(v)}' 해석 불가")
                continue
            if isinstance(v, str) and v == t:
                continue
            if _display_value(v) == t:
                row[col] = t
                silent_count += 1
            else:
                job.proposals.append(FixProposal(
                    "시간형식", label, rlabel, person,
                    _display_value(v), t, f"{colname}을 4자리 형식으로 통일",
                    [(i, col)], _single_time_builder,
                ))

        for col, colname, builder in (
            (COL_DATE, "서비스일자", _date_builder),
            (COL_BIRTH, "대상자 생년월일", _date_builder),
        ):
            v = row[col]
            if _is_blank(v):
                continue
            d = normalize_date_value(v)
            if d is None:
                continue
            if isinstance(v, int) and v == d:
                continue
            if _display_value(v) == str(d):
                row[col] = d
                silent_count += 1
            else:
                job.proposals.append(FixProposal(
                    "시간형식", label, rlabel, person,
                    _display_value(v), str(d), f"{colname}을 숫자 형식으로 통일",
                    [(i, col)], builder,
                ))

        v = row[COL_CODE]
        cleaned = clean_service_code(v)
        if not _is_blank(v) and cleaned != ("" if v is None else str(v)):
            job.proposals.append(FixProposal(
                "코드정리", label, rlabel, person,
                repr(_display_value(v)), cleaned, "서비스코드의 공백/특수문자 제거",
                [(i, COL_CODE)], _code_builder,
            ))

        vv = row[COL_VENDOR]
        vendor_fixed = normalize_vendor_code(vv)
        if _is_blank(vv):
            row[COL_VENDOR] = vendor_fixed
            job.auto_fixes.append(f"{rlabel}행 거래처코드가 비어있어 기본값 '{vendor_fixed}'를 채워넣었습니다")
        elif vendor_fixed != _norm_str(vv):
            old_disp = _display_value(vv)
            row[COL_VENDOR] = vendor_fixed
            job.auto_fixes.append(f"{rlabel}행 거래처코드 '{old_disp}'→'{vendor_fixed}' 자동 보정 (4자리로 통일)")

        if _is_blank(row[COL_WORKER]):
            job.errors.append(f"{rlabel}행 생활지원사명이 비어있음")

        if not cleaned:
            job.errors.append(f"{rlabel}행 서비스코드 없음")
        elif cleaned not in SERVICE_CODES:
            job.errors.append(f"{rlabel}행 서비스코드 '{cleaned}' 코드표에 없음")

        if cleaned == "1A102":
            s = normalize_time_value(row[COL_START])
            e = normalize_time_value(row[COL_END])
            if s is not None and e is not None:
                duration = time_to_minutes(e) - time_to_minutes(s)
                if 0 < duration <= 9:
                    job.dropped_rows[i] = (
                        f"{rlabel}행 {person} 서비스코드 1A102(방문)인데 계획시간이 {duration}분으로 9분 이하임")

    if silent_count:
        job.auto_fixes.append(f"표시가 같은 저장 형식 통일 {silent_count}건 자동 적용")

    def _to_date(n):
        try:
            return datetime.date(n // 10000, (n // 100) % 100, n % 100)
        except ValueError:
            return None

    parsed_dates = {}
    for i, row in enumerate(grid):
        n = normalize_date_value(row[COL_DATE])
        if n is not None:
            parsed_dates[i] = (n, _to_date(n))
    if parsed_dates:
        current_year = datetime.date.today().year

        for i, (n, dt) in list(parsed_dates.items()):
            year = n // 10000
            if year == current_year:
                continue
            mm, dd = (n // 100) % 100, n % 100
            fixed = current_year * 10000 + mm * 100 + dd
            fixed_dt = _to_date(fixed)
            rlabel = row_nums[i]
            person = _norm_str(grid[i][COL_NAME])
            if fixed_dt is None:
                alt = fixed - 1
                alt_dt = _to_date(alt)
                if alt_dt is not None:
                    fixed, fixed_dt = alt, alt_dt
                    why = f"연도가 올해({current_year}년)가 아니고, 올해는 {mm}월 {dd}일이 없어 {mm}월 {dd - 1}일로 자동 보정 ({year}년→{current_year}년)"
                else:
                    job.warnings.append(
                        f"{rlabel}행 서비스일자 {n}의 연도가 올해가 아닌데, 올해 달력에 맞는 날짜로 자동 보정하지 못했습니다 (직접 확인 필요)")
                    continue
            else:
                why = f"연도가 올해({current_year}년)가 아니어서 자동 보정 ({year}년→{current_year}년)"
            job.proposals.append(FixProposal(
                "날짜", label, rlabel, person,
                str(n), str(fixed), why,
                [(i, COL_DATE)], _date_builder,
            ))
            parsed_dates[i] = (fixed, fixed_dt)

        major_dates = sorted(
            dt for n, dt in parsed_dates.values()
            if dt is not None and n // 10000 == current_year
        )
        margin = datetime.timedelta(days=7)
        med = major_dates[len(major_dates) // 2] if major_dates else None

        for i, (n, dt) in parsed_dates.items():
            if n // 10000 != current_year:
                continue
            rlabel = row_nums[i]
            person = _norm_str(grid[i][COL_NAME])
            if dt is None:
                yy, mm0, dd0 = n // 10000, (n // 100) % 100, n % 100
                last_day = calendar.monthrange(yy, mm0)[1]
                job.dropped_rows[i] = (
                    f"{rlabel}행 {person} 서비스일자 {yy}-{mm0:02d}-{dd0:02d}는 달력에 없는 날짜임 "
                    f"({mm0}월은 {last_day}일까지)")
                continue
            if med is not None and abs(dt - med) <= margin:
                continue
            mm, dd = (n // 100) % 100, n % 100
            fixed = None
            why = ""
            if med is not None and dd <= 12:
                cand = _to_date(current_year * 10000 + dd * 100 + mm)
                if cand is not None and abs(cand - med) <= margin:
                    fixed = current_year * 10000 + dd * 100 + mm
                    why = "월/일이 뒤바뀐 것으로 추정"
            if fixed is not None and fixed != n:
                job.proposals.append(FixProposal(
                    "날짜", label, rlabel, person,
                    str(n), str(fixed), why,
                    [(i, COL_DATE)], _date_builder,
                ))
            else:
                job.warnings.append(
                    f"{rlabel}행 서비스일자 {n}가 파일의 다른 날짜들과 동떨어져 있음")

    events_by_key = {}
    for i, row in enumerate(grid):
        d = normalize_date_value(row[COL_DATE])
        s = normalize_time_value(row[COL_START])
        e = normalize_time_value(row[COL_END])
        name = _norm_str(row[COL_NAME])
        if d is None or s is None or e is None or not name or i in job.dropped_rows:
            continue
        if time_to_minutes(e) <= time_to_minutes(s):
            job.dropped_rows[i] = (
                f"{row_nums[i]}행 {name} 종료시간({e})이 시작시간({s})보다 빠르거나 같음 (시간이 음수가 됨)")
            continue
        events_by_key.setdefault((name, d), []).append(
            (time_to_minutes(s), time_to_minutes(e), i))

    for (name, d), events in events_by_key.items():
        events.sort()
        for k in range(len(events) - 1):
            s1, e1, i1 = events[k]
            s2, e2, i2 = events[k + 1]
            if s2 >= e1:
                continue
            rlabel1, rlabel2 = job.row_nums[i1], job.row_nums[i2]
            date_disp = f"{(d % 10000) // 100}월{d % 100}일"
            if e2 >= e1:
                job.proposals.append(FixProposal(
                    "시간겹침", label, rlabel1, name,
                    minutes_to_time(e1), minutes_to_time(s2),
                    f"{date_disp} {minutes_to_time(s1)}~{minutes_to_time(e1)}와 "
                    f"{minutes_to_time(s2)}~{minutes_to_time(e2)}({rlabel2}행)이 겹쳐 종료시간을 당김",
                    [(i1, COL_END)], _single_time_builder,
                ))
                events[k] = (s1, s2, i1)
            else:
                dur = e2 - s2
                new_s, new_e = e1, e1 + dur
                job.proposals.append(FixProposal(
                    "시간겹침", label, rlabel2, name,
                    f"{minutes_to_time(s2)}~{minutes_to_time(e2)}",
                    f"{minutes_to_time(new_s)}~{minutes_to_time(new_e)}",
                    f"{date_disp} 일정이 {minutes_to_time(s1)}~{minutes_to_time(e1)}({rlabel1}행) "
                    f"안에 완전히 포함되어 뒤로 이동",
                    [(i2, COL_START), (i2, COL_END)], _time_range_builder,
                ))
                events[k + 1] = (new_s, new_e, i2)
                events.sort()


def _cross_check_beneficiaries(job: FileJob, beneficiaries):
    grid, row_nums = job.grid, job.row_nums
    for i, row in enumerate(grid):
        raw_name = _norm_str(row[COL_NAME])
        if not raw_name:
            continue
        rlabel = row_nums[i]
        m = re.match(r"^([가-힣]{2,})([A-Za-z])?$", raw_name)
        base_name = m.group(1) if m else raw_name
        birth = normalize_date_value(row[COL_BIRTH])
        gender = _norm_str(row[COL_GENDER])
        worker = _norm_str(row[COL_WORKER])

        entries = beneficiaries.get(raw_name) or beneficiaries.get(base_name)
        if not entries:
            candidates = [
                e for group in beneficiaries.values() for e in group
                if birth is not None and e["birth"] == birth and e["worker"] == worker
            ]
            if len(candidates) == 1:
                fixed = candidates[0].get("display_name", candidates[0]["name"])
                if fixed != raw_name:
                    grid[i][COL_NAME] = fixed
                    job.auto_fixes.append(
                        f"{rlabel}행 대상자명 '{raw_name}'→'{fixed}' (리스트 대조: 생년월일·지원사 일치)")
                _check_group_service_code(job, rlabel, fixed, candidates[0], row)
            else:
                job.errors.append(f"{rlabel}행 대상자 '{raw_name}' 대상자리스트에 없음")
            continue

        matched = [e for e in entries if birth is not None and e["birth"] == birth]
        if len(matched) != 1:
            by_worker = [e for e in entries if e["worker"] == worker]
            if len(by_worker) == 1:
                matched = by_worker
            elif len(entries) == 1:
                matched = entries
            else:
                job.warnings.append(
                    f"{rlabel}행 대상자 '{raw_name}' 동명이인 구분 불가 (생년월일 불일치)")
                continue

        entry = matched[0]
        display_name = entry.get("display_name", entry["name"])
        if display_name != raw_name:
            grid[i][COL_NAME] = display_name
            job.auto_fixes.append(
                f"{rlabel}행 대상자명 '{raw_name}'→'{display_name}' (생년월일로 동명이인 구분)")
        if birth != entry["birth"] and entry["birth"] is not None:
            old_birth = _display_value(row[COL_BIRTH]) if birth is None else str(birth)
            grid[i][COL_BIRTH] = entry["birth"]
            job.auto_fixes.append(
                f"{rlabel}행 {raw_name} 생년월일 {old_birth}→{entry['birth']} (대상자리스트 기준)")
        if gender and entry["gender"] and gender != entry["gender"]:
            grid[i][COL_GENDER] = entry["gender"]
            job.auto_fixes.append(
                f"{rlabel}행 {raw_name} 성별 {gender}→{entry['gender']} (대상자리스트 기준)")
        elif not gender and entry["gender"]:
            grid[i][COL_GENDER] = entry["gender"]
            job.auto_fixes.append(f"{rlabel}행 {raw_name} 성별 빈칸→{entry['gender']} (대상자리스트 기준)")

        if worker and entry["worker"] and worker != entry["worker"]:
            job.warnings.append(
                f"{rlabel}행 대상자 '{raw_name}'의 담당 지원사가 리스트({entry['worker']})와 다름({worker})")

        _check_group_service_code(job, rlabel, display_name, entry, row)


def _check_group_service_code(job: FileJob, rlabel, name, entry, row):
    if not entry.get("group", "").startswith("일반"):
        return
    code = clean_service_code(row[COL_CODE])
    if code in GENERAL_GROUP_FORBIDDEN_CODES:
        job.errors.append(
            f"{rlabel}행 {name} 일반군인데 서비스코드 {code}({GENERAL_GROUP_FORBIDDEN_CODES[code]})가 들어감")


def _cross_check_workers(job: FileJob, worker_data):
    grid, row_nums = job.grid, job.row_nums
    for i, row in enumerate(grid):
        name = _norm_str(row[COL_WORKER])
        if not name:
            continue
        rlabel = row_nums[i]
        known = worker_data.get(name)
        if known is None:
            job.warnings.append(
                f"{rlabel}행 생활지원사 '{name}'이(가) 등록된 지원사 명단에 없음 (오타 또는 신규 지원사 확인 필요)")
            continue

        cur_birth = normalize_date_value(row[COL_WORKER_BIRTH])
        known_birth = known["birth"] if isinstance(known, dict) else known
        if cur_birth != known_birth:
            old_disp = _display_value(row[COL_WORKER_BIRTH])
            grid[i][COL_WORKER_BIRTH] = known_birth
            job.auto_fixes.append(
                f"{rlabel}행 {name} 생활지원사 생년월일 {old_disp}→{known_birth} (지원사 명단 기준)")


def send_to_recycle_bin(paths):
    import ctypes

    class SHFILEOPSTRUCTW(ctypes.Structure):
        _fields_ = [
            ("hwnd", ctypes.c_void_p),
            ("wFunc", ctypes.c_uint),
            ("pFrom", ctypes.c_wchar_p),
            ("pTo", ctypes.c_wchar_p),
            ("fFlags", ctypes.c_ushort),
            ("fAnyOperationsAborted", ctypes.c_int),
            ("hNameMappings", ctypes.c_void_p),
            ("lpszProgressTitle", ctypes.c_wchar_p),
        ]

    FO_DELETE = 3
    FOF_ALLOWUNDO = 0x40
    FOF_NOCONFIRMATION = 0x10
    FOF_SILENT = 0x0004

    paths = [str(p) for p in paths]
    if not paths:
        return True
    op = SHFILEOPSTRUCTW()
    op.wFunc = FO_DELETE
    op.pFrom = "\0".join(paths) + "\0"
    op.fFlags = FOF_ALLOWUNDO | FOF_NOCONFIRMATION | FOF_SILENT
    result = ctypes.windll.shell32.SHFileOperationW(ctypes.byref(op))
    return result == 0 and not op.fAnyOperationsAborted


MERGED_FILE_BASE_NAME = "★통합본★"
ERROR_FILE_PREFIX = "오류("


def merge_converted_files(target_folder: Path, log=None):
    def _log(msg):
        if log:
            log(msg)

    files = [
        f for f in target_folder.glob("*.xlsx")
        if not f.stem.startswith(MERGED_FILE_BASE_NAME)
        and not f.stem.startswith(ERROR_FILE_PREFIX)
        and not f.name.startswith("~$")
    ]
    if not files:
        return None
    files.sort(key=_natural_sort_key)

    merged_wb = openpyxl.Workbook()
    merged_ws = merged_wb.active

    out_row = 1
    header_written = False

    for f in files:
        wb = openpyxl.load_workbook(_read_bytes_io(f), data_only=True)
        ws = wb.active
        max_col = ws.max_column

        if not header_written:
            for col in range(1, max_col + 1):
                merged_ws.cell(row=out_row, column=col, value=ws.cell(row=1, column=col).value)
            out_row += 1
            header_written = True

        for row in range(2, ws.max_row + 1):
            row_values = [ws.cell(row=row, column=col).value for col in range(1, max_col + 1)]
            if all(v is None or (isinstance(v, str) and v.strip() == "") for v in row_values):
                continue
            for col, value in enumerate(row_values, start=1):
                merged_ws.cell(row=out_row, column=col, value=value)
            out_row += 1

    out_path = _unique_target_path(target_folder, MERGED_FILE_BASE_NAME)
    merged_wb.save(out_path)
    _log(f"{out_path.name} 저장 완료 (총 {out_row - 1}행)")
    return out_path


def _data_date_range(job: FileJob):
    dates = [normalize_date_value(r[COL_DATE]) for r in job.grid]
    dates = [d for d in dates if d is not None]
    if not dates:
        return None, None
    lo, hi = min(dates), max(dates)
    return ((lo % 10000) // 100, lo % 100), ((hi % 10000) // 100, hi % 100)


def _output_base_name(job: FileJob) -> str:
    name = job.worker_name

    start, end = job.file_date_start, job.file_date_end
    if start is None:
        start, end = _data_date_range(job)
        if end == start:
            end = None

    date_part = ""
    if start:
        date_part = f" {start[0]}월{start[1]}일"
        if end and end != start:
            date_part += f"-{end[0]}월{end[1]}일"
    rev_part = "(수정)" if job.is_revision else ""

    if job.errors:
        return f"{ERROR_FILE_PREFIX}{name}){date_part}{rev_part}"
    return f"{name}{date_part}{rev_part}"


def _check_filename_vs_data(job: FileJob, log):
    if job.file_date_start is None:
        log(f"  [확인필요] 파일명에서 날짜를 읽지 못했습니다: {job.src_path.name}")
        return
    data_start, data_end = _data_date_range(job)
    if data_start is None:
        return
    if data_start != job.file_date_start:
        log(f"  [확인필요] 파일명 시작일({job.file_date_start[0]}월{job.file_date_start[1]}일)과 "
            f"데이터 시작일({data_start[0]}월{data_start[1]}일)이 다릅니다.")
    if job.file_date_end is not None and data_end != job.file_date_end:
        log(f"  [확인필요] 파일명 종료일({job.file_date_end[0]}월{job.file_date_end[1]}일)과 "
            f"데이터 마지막 날짜({data_end[0]}월{data_end[1]}일)가 다릅니다.")


def auto_approve_all(proposals):
    return {p.pid: p.suggested_disp for p in proposals}


def _split_row(msg):
    m = re.match(r"\s*(\d+)행\s*(.*)", msg)
    if m:
        return int(m.group(1)), m.group(2)
    return None, msg


def _log_file_block(log, tag, name, messages):
    items = sorted((_split_row(m) for m in messages),
                   key=lambda x: (x[0] is None, x[0] or 0))
    log(f"{tag} {name} ({len(items)}건)")
    for row, text in items:
        if row is None:
            log(f"          {text}")
        else:
            log(f"    {row:>3}행   {text}")


def _log_final_summary(jobs, log):
    error_jobs = [j for j in jobs if j.errors]
    dropped_jobs = [j for j in jobs if j.dropped_rows]
    if error_jobs:
        log("")
        log(f"오류 파일 {len(error_jobs)}개는 통합본에서 제외합니다:")
        for j in error_jobs:
            _log_file_block(log, "[오류]", j.worker_name, j.errors)
    if dropped_jobs:
        log("")
        log(f"잘못된 줄을 삭제하고 저장한 파일 {len(dropped_jobs)}개:")
        for j in dropped_jobs:
            _log_file_block(log, "[줄 삭제]", j.worker_name, list(j.dropped_rows.values()))


def run_conversion(log, roster_path, worker_data=None, on_progress=None, review_callback=None):
    """roster_path: 오늘 대상자리스트 파일. worker_data: {지원사 이름: {"birth": YYYYMMDD}} (없으면 지원사 검증 생략)"""
    source_folder, target_folder = work_folders()
    beneficiaries = load_beneficiary_list(roster_path, log)  # 못 읽으면 변환하지 않고 멈춤
    n = sum(len(v) for v in beneficiaries.values())
    log(f"대상자리스트 {n}명 읽기 완료 ({Path(roster_path).name})")
    if worker_data:
        log(f"생활지원사 명단 {len(worker_data)}명 (goodeos)")
    else:
        log("[경고] 생활지원사 명단이 없어 지원사 생년월일 검증은 건너뜁니다.")

    src_files = find_source_excels(source_folder)
    total = len(src_files)
    log(f"대상 파일 {total}개 발견")
    if on_progress:
        on_progress(0, total)

    jobs = []
    failed_files = []
    for i, src_file in enumerate(src_files):
        log(f"검사 중: {src_file.name}")
        job = FileJob(src_file)
        try:
            wb = load_workbook_safely(src_file, log)
            scan_file(job, wb.active, beneficiaries, log, worker_data)
            if job.load_error:
                log(f"  -> {job.load_error}, 건너뜁니다.")
                failed_files.append((src_file.name, job.load_error))
            else:
                jobs.append(job)
        except Exception as e:
            log(f"  -> [실패] {e}")
            failed_files.append((src_file.name, str(e)))
        finally:
            if on_progress:
                on_progress(i + 1, total)

    all_proposals = [p for job in jobs for p in job.proposals]
    approved = {}
    if all_proposals:
        log("")
        log(f"승인이 필요한 수정 제안 {len(all_proposals)}건 발견 - 검토 창을 확인하세요.")
        if review_callback is None:
            review_callback = auto_approve_all
        approved = review_callback(all_proposals) or {}

    success_count = 0
    applied_count = 0
    skipped_count = 0
    deletable_sources = []
    deletable_outputs = []
    for job in jobs:
        log(f"저장 중: {job.src_path.name}")

        for msg in job.auto_fixes:
            log(f"  [자동수정] {msg}")
        for msg in job.warnings:
            log(f"  [확인필요] {msg}")

        for p in job.proposals:
            if p.pid not in approved:
                skipped_count += 1
                log(f"  [건너뜀] {p.row_label}행 {p.kind}: {p.current_disp} → {p.suggested_disp}")
                continue
            try:
                values = p.build_values(approved[p.pid])
            except ValueError as e:
                skipped_count += 1
                log(f"  [건너뜀] {p.row_label}행 {p.kind}: 입력값 오류 ({e})")
                continue
            for (r, c), v in zip(p.cells, values):
                job.grid[r][c] = v
            applied_count += 1
            log(f"  [수정] {p.row_label}행 {p.kind}: {p.current_disp} → {approved[p.pid]}")

        for msg in job.errors:
            log(f"  [오류] {msg}")
        for i in sorted(job.dropped_rows):
            log(f"  [줄 삭제] {job.dropped_rows[i]} → 이 줄을 빼고 저장합니다")

        _check_filename_vs_data(job, log)

        try:
            new_wb = openpyxl.Workbook()
            new_ws = new_wb.active
            for col, header in enumerate(STANDARD_HEADERS, start=1):
                new_ws.cell(row=1, column=col, value=header)
            kept_rows = [row for i, row in enumerate(job.grid) if i not in job.dropped_rows]
            for r, row in enumerate(kept_rows, start=2):
                for c, value in enumerate(row, start=1):
                    new_ws.cell(row=r, column=c, value=value)

            out_path = _unique_target_path(target_folder, _output_base_name(job))
            new_wb.save(out_path)
            if job.errors:
                log(f"  -> [{out_path.stem}] 저장 완료 (오류 {len(job.errors)}건 - 통합본에서 제외)")
            else:
                log(f"  -> [{out_path.stem}] 저장 완료")
                deletable_sources.append(job.src_path)
                deletable_outputs.append(out_path)
            success_count += 1
        except Exception as e:
            log(f"  -> [실패] {e}")
            failed_files.append((job.src_path.name, str(e)))

    if all_proposals:
        log("")
        log(f"수정 제안 처리 결과: 적용 {applied_count}건 / 건너뜀 {skipped_count}건")

    try:
        merged_path = merge_converted_files(target_folder, log)
        if merged_path is not None:
            deletable_outputs.append(merged_path)
    except Exception as e:
        log(f"통합본 생성 중 오류: {e}")

    _log_final_summary(jobs, log)

    return target_folder, total, success_count, failed_files, deletable_sources, deletable_outputs
